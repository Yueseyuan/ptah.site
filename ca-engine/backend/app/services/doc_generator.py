import os
import json
from datetime import datetime
from typing import Dict, Any, List, Union
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from app.config import settings

DOCS_DIR = "generated_docs"
os.makedirs(DOCS_DIR, exist_ok=True)

NAVY  = RGBColor(0x0d, 0x22, 0x44)
GOLD  = RGBColor(0xc4, 0x9a, 0x1a)
BLACK = RGBColor(0x1a, 0x25, 0x35)

def _filename(prefix: str, client_id: Any) -> str:
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    return f"{DOCS_DIR}/{prefix}_{client_id}_{ts}.docx"

def _base_doc(title: str) -> Document:
    doc = Document()
    # Page margins
    for section in doc.sections:
        section.top_margin    = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin   = Inches(1.25)
        section.right_margin  = Inches(1.25)

    # Header
    hdr = doc.sections[0].header
    hp  = hdr.paragraphs[0]
    hp.text = f"{settings.COMPANY_NAME}  ·  {settings.COMPANY_PHONE}  ·  {settings.COMPANY_EMAIL}"
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in hp.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = GOLD

    # Title block
    tp = doc.add_paragraph()
    tr = tp.add_run(title.upper())
    tr.bold = True
    tr.font.size = Pt(14)
    tr.font.color.rgb = NAVY
    tp.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph(f"Date: {datetime.now().strftime('%B %d, %Y')}").paragraph_format.space_after = Pt(6)
    doc.add_paragraph()
    return doc

def _sig_block(doc: Document, label: str = "Client", include_date: bool = True):
    doc.add_paragraph()
    doc.add_paragraph(f"{'_' * 45}     {'_' * 20}")
    sig_line = doc.add_paragraph()
    sig_line.add_run(f"{label} Signature").bold = True
    sig_line.add_run("          Date" if include_date else "")
    doc.add_paragraph()
    doc.add_paragraph(f"{'_' * 45}")
    doc.add_paragraph(f"Printed Name")

def _add_disclosure(doc: Document):
    doc.add_paragraph()
    p = doc.add_paragraph()
    r = p.add_run("NON-ATTORNEY DISCLOSURE: ")
    r.bold = True
    r.font.size = Pt(9)
    p.add_run(
        f"{settings.COMPANY_NAME} is not a law firm and does not provide legal advice or "
        "legal representation. All services are administrative and educational in nature only. "
        "Any legal services require referral to a licensed attorney."
    ).font.size = Pt(9)


# ─────────────────────────────────────────────
# DIVISION 1 — NOTARY & SIGNING
# ─────────────────────────────────────────────

def gen_notary_service_agreement(data: Dict[str, Any]) -> str:
    doc = _base_doc("Notary Service Agreement")
    name = f"{data.get('first_name','')} {data.get('last_name','')}".strip()

    doc.add_paragraph(f"Client: {name}")
    doc.add_paragraph(f"Address: {data.get('address','')}, {data.get('city','')}, {data.get('state','SC')} {data.get('zip_code','')}")
    doc.add_paragraph(f"Phone: {data.get('phone','')}")
    doc.add_paragraph(f"Email: {data.get('email','')}")
    doc.add_paragraph(f"Appointment Date/Time: {data.get('appointment_datetime','')}")
    doc.add_paragraph(f"Location: {data.get('location','')}")
    doc.add_paragraph()

    doc.add_paragraph("SERVICES REQUESTED:", style="Heading 2")
    doc.add_paragraph(f"Document Type: {data.get('document_type','')}")
    doc.add_paragraph(f"Number of Signers: {data.get('num_signers', 1)}")
    doc.add_paragraph(f"Number of Notarial Acts: {data.get('num_acts', 1)}")
    doc.add_paragraph()

    doc.add_paragraph("FEE SCHEDULE (South Carolina):", style="Heading 2")
    num_acts = int(data.get('num_acts', 1))
    act_fee  = 5.00
    travel   = float(data.get('travel_fee', 0))
    svc_fee  = float(data.get('service_fee', 0))
    total    = (num_acts * act_fee) + travel + svc_fee

    doc.add_paragraph(f"Notarial Acts ({num_acts} × $5.00 per SC Code §26-1-120): ${num_acts * act_fee:.2f}")
    if travel:
        doc.add_paragraph(f"Travel Fee: ${travel:.2f}")
    if svc_fee:
        doc.add_paragraph(f"Signing Agent Service Fee: ${svc_fee:.2f}")
    doc.add_paragraph(f"TOTAL: ${total:.2f}")
    doc.add_paragraph()

    doc.add_paragraph("TERMS:", style="Heading 2")
    doc.add_paragraph(
        "Payment is due at time of service. Cancellations within 24 hours of appointment may incur a "
        "cancellation fee. Client agrees to present valid government-issued photo ID at time of service."
    )
    _add_disclosure(doc)
    doc.add_paragraph()
    _sig_block(doc, "Client")
    doc.add_paragraph()
    _sig_block(doc, f"Notary — {settings.COMPANY_NAME}")

    path = _filename("notary_agreement", data.get('client_id','new'))
    doc.save(path)
    return path


def gen_notary_invoice(data: Dict[str, Any]) -> str:
    doc = _base_doc(f"Invoice #{data.get('invoice_number', 'CA-' + datetime.now().strftime('%Y%m%d'))}")
    name = f"{data.get('first_name','')} {data.get('last_name','')}".strip()

    doc.add_paragraph(f"Bill To: {name}")
    doc.add_paragraph(f"Phone: {data.get('phone','')}")
    doc.add_paragraph(f"Date of Service: {data.get('service_date', datetime.now().strftime('%B %d, %Y'))}")
    doc.add_paragraph()

    # Table
    tbl = doc.add_table(rows=1, cols=4)
    tbl.style = 'Table Grid'
    for i, h in enumerate(['Service', 'Qty', 'Unit Price', 'Amount']):
        tbl.rows[0].cells[i].text = h

    services = data.get('services', [{'description': 'Notary Service', 'quantity': 1, 'unit_price': 5.0}])
    total = 0.0
    for svc in services:
        row = tbl.add_row().cells
        amt = float(svc.get('quantity', 1)) * float(svc.get('unit_price', 0))
        row[0].text = svc.get('description', '')
        row[1].text = str(svc.get('quantity', 1))
        row[2].text = f"${float(svc.get('unit_price', 0)):.2f}"
        row[3].text = f"${amt:.2f}"
        total += amt

    doc.add_paragraph()
    doc.add_paragraph(f"TOTAL DUE: ${total:.2f}")
    doc.add_paragraph(f"Payment Methods: Cash · Zelle · CashApp · Credit Card")
    doc.add_paragraph(f"Thank you for choosing {settings.COMPANY_NAME}!")

    path = _filename("invoice", data.get('client_id', 'new'))
    doc.save(path)
    return path


# ─────────────────────────────────────────────
# DIVISION 2 — CREDIT RESTORATION
# ─────────────────────────────────────────────

def gen_credit_services_agreement(data: Dict[str, Any]) -> str:
    doc = _base_doc("Credit Services Organization Agreement")
    name = f"{data.get('first_name','')} {data.get('last_name','')}".strip()

    doc.add_paragraph(f"Client: {name}")
    doc.add_paragraph(f"Address: {data.get('address','')}, {data.get('city','')}, {data.get('state','SC')} {data.get('zip_code','')}")
    doc.add_paragraph(f"Email: {data.get('email','')}")
    doc.add_paragraph(f"Phone: {data.get('phone','')}")
    doc.add_paragraph()

    doc.add_paragraph("SERVICES SELECTED:", style="Heading 2")
    for svc in data.get('selected_services', ['Credit Investigation Package']):
        doc.add_paragraph(f"  • {svc}")
    doc.add_paragraph(f"\nTotal Fee: ${data.get('total_fee', 0):.2f}")
    doc.add_paragraph(f"Payment Plan: {data.get('payment_plan', 'Paid in full')}")
    doc.add_paragraph()

    doc.add_paragraph("CREDIT SERVICES DISCLOSURE (Required by Law):", style="Heading 2")
    doc.add_paragraph(
        "You have a right to dispute inaccurate information in your credit report without charge. "
        "You may contact the consumer reporting agency directly. This service is not legal advice. "
        "Results are not guaranteed. You may cancel this agreement within 3 business days without penalty."
    )

    _add_disclosure(doc)
    doc.add_paragraph()
    _sig_block(doc, "Client")
    doc.add_paragraph()
    _sig_block(doc, f"{settings.COMPANY_NAME} Representative")

    path = _filename("credit_agreement", data.get('client_id', 'new'))
    doc.save(path)
    return path


def gen_dispute_letter(data: Dict[str, Any]) -> str:
    bureau = data.get('bureau', 'Credit Bureau')
    bureau_addresses = {
        'Equifax':    'Equifax Information Services LLC, P.O. Box 740256, Atlanta, GA 30374',
        'Experian':   'Experian, P.O. Box 4500, Allen, TX 75013',
        'TransUnion': 'TransUnion LLC Consumer Dispute Center, P.O. Box 2000, Chester, PA 19016',
    }
    address = bureau_addresses.get(bureau, f'{bureau}, [Address]')

    doc = _base_doc(f"Credit Dispute Letter — {bureau}")
    name = f"{data.get('first_name','')} {data.get('last_name','')}".strip()

    doc.add_paragraph(f"To: {address}")
    doc.add_paragraph()
    doc.add_paragraph(f"Re: Formal Dispute — {name}")
    doc.add_paragraph(f"SSN (Last 4): XXX-XX-{data.get('ssn_last4','****')}")
    doc.add_paragraph(f"Date of Birth: {data.get('dob','')}")
    doc.add_paragraph(f"Current Address: {data.get('address','')}, {data.get('city','')}, {data.get('state','SC')} {data.get('zip_code','')}")
    doc.add_paragraph()

    doc.add_paragraph(
        f"I am writing pursuant to the Fair Credit Reporting Act (15 U.S.C. § 1681 et seq.) to formally "
        f"dispute the following items appearing on my credit report maintained by {bureau}. "
        "I request that you investigate and correct or remove each item below within 30 days."
    )
    doc.add_paragraph()

    for i, item in enumerate(data.get('dispute_items', []), 1):
        doc.add_paragraph(f"Item {i}:", style="Heading 3")
        doc.add_paragraph(f"  Creditor/Account: {item.get('creditor','')}")
        doc.add_paragraph(f"  Account Number: {item.get('account_number','')}")
        doc.add_paragraph(f"  Reason: {item.get('reason','')}")
        doc.add_paragraph(f"  Requested Action: {item.get('action', 'Remove or correct this item.')}")
        doc.add_paragraph()

    doc.add_paragraph(
        "If you cannot verify the accuracy of any disputed item, you must delete it from my report. "
        "Please provide me with written results of your investigation and a free copy of my updated report."
    )
    doc.add_paragraph()
    doc.add_paragraph("Enclosures: Copy of valid photo ID · Proof of address · Supporting documentation")
    doc.add_paragraph()
    doc.add_paragraph(f"Sincerely,\n\n{name}")
    doc.add_paragraph(f"Phone: {data.get('phone','')}")
    doc.add_paragraph(f"Email: {data.get('email','')}")

    path = _filename(f"dispute_{bureau.lower()}", data.get('client_id', 'new'))
    doc.save(path)
    return path


def gen_debt_validation_letter(data: Dict[str, Any]) -> str:
    doc = _base_doc("Debt Validation Letter")
    name = f"{data.get('first_name','')} {data.get('last_name','')}".strip()

    doc.add_paragraph(f"To: {data.get('collector_name','Debt Collector')}")
    doc.add_paragraph(f"Address: {data.get('collector_address','')}")
    doc.add_paragraph()
    doc.add_paragraph(f"Re: Debt Validation Request — Account #{data.get('account_number','')}")
    doc.add_paragraph()

    doc.add_paragraph(
        f"I, {name}, hereby request validation of the alleged debt referenced above pursuant to "
        "the Fair Debt Collection Practices Act (15 U.S.C. § 1692g). "
        "Please provide the following within 30 days:"
    )
    doc.add_paragraph()
    for item in [
        "Complete account history and original signed agreement",
        "Name and address of original creditor",
        "Proof you are licensed to collect in South Carolina",
        "Copy of last billing statement from original creditor",
        "Verification that the statute of limitations has not expired",
        "Breakdown of all fees, interest, and principal claimed",
    ]:
        doc.add_paragraph(f"  • {item}")

    doc.add_paragraph()
    doc.add_paragraph(
        "Until validation is provided, please cease all collection activity. "
        "This is not a refusal to pay, but a formal request for validation as is my right under federal law."
    )
    doc.add_paragraph()
    doc.add_paragraph(f"Sincerely,\n\n{name}")

    path = _filename("debt_validation", data.get('client_id', 'new'))
    doc.save(path)
    return path


def gen_goodwill_letter(data: Dict[str, Any]) -> str:
    doc = _base_doc("Goodwill Letter")
    name = f"{data.get('first_name','')} {data.get('last_name','')}".strip()

    doc.add_paragraph(f"To: {data.get('creditor_name','Creditor')} Customer Relations")
    doc.add_paragraph()

    doc.add_paragraph(
        f"My name is {name} and I have been a customer since {data.get('account_open_date','')}. "
        f"I am writing to respectfully request a goodwill adjustment to remove a late payment recorded "
        f"on {data.get('late_payment_date','')} from my credit report."
    )
    doc.add_paragraph()
    doc.add_paragraph(
        f"The late payment occurred due to {data.get('hardship_reason','an unexpected financial hardship')}. "
        "This was an isolated incident and does not reflect my typical payment behavior. "
        "Since that time, my account has been in good standing."
    )
    doc.add_paragraph()
    doc.add_paragraph(
        "I have been a loyal customer and this negative mark is significantly impacting my ability to "
        "achieve my financial goals. I respectfully ask that you consider removing this entry as a "
        "gesture of goodwill. I am deeply committed to maintaining my financial responsibilities."
    )
    doc.add_paragraph()
    doc.add_paragraph(f"Thank you sincerely for your consideration.\n\n{name}")
    doc.add_paragraph(f"Account Number: {data.get('account_number','')}")
    doc.add_paragraph(f"Phone: {data.get('phone','')}")

    path = _filename("goodwill_letter", data.get('client_id', 'new'))
    doc.save(path)
    return path


# ─────────────────────────────────────────────
# DIVISION 3 — CRIMINAL RECORD RELIEF & REENTRY
# ─────────────────────────────────────────────

def gen_reentry_service_agreement(data: Dict[str, Any]) -> str:
    doc = _base_doc("Criminal Record Relief & Reentry — Service Agreement")
    name = f"{data.get('first_name','')} {data.get('last_name','')}".strip()

    doc.add_paragraph(f"Client: {name}")
    doc.add_paragraph(f"Address: {data.get('address','')}, {data.get('city','')}, {data.get('state','SC')} {data.get('zip_code','')}")
    doc.add_paragraph(f"Phone: {data.get('phone','')}")
    doc.add_paragraph(f"Email: {data.get('email','')}")
    doc.add_paragraph()

    doc.add_paragraph("ADMINISTRATIVE SERVICES REQUESTED:", style="Heading 2")
    for svc in data.get('selected_services', []):
        doc.add_paragraph(f"  • {svc}")
    doc.add_paragraph()

    doc.add_paragraph("IMPORTANT — SCOPE OF SERVICES:", style="Heading 2")
    doc.add_paragraph(
        f"{settings.COMPANY_NAME} provides ADMINISTRATIVE ASSISTANCE ONLY. Services include record "
        "retrieval, file organization, packet assembly, and referrals. We do NOT provide legal advice, "
        "file court documents, or represent clients. For legal representation, clients will be referred "
        "to licensed attorneys."
    )

    doc.add_paragraph(f"\nService Fee: ${data.get('service_fee', 0):.2f}")
    _add_disclosure(doc)
    doc.add_paragraph()
    _sig_block(doc, "Client")
    doc.add_paragraph()
    _sig_block(doc, f"{settings.COMPANY_NAME} Representative")

    path = _filename("reentry_agreement", data.get('client_id', 'new'))
    doc.save(path)
    return path


def gen_non_attorney_disclosure(data: Dict[str, Any]) -> str:
    doc = _base_doc("Non-Attorney Disclosure & Client Acknowledgment")
    name = f"{data.get('first_name','')} {data.get('last_name','')}".strip()

    doc.add_paragraph(f"Client: {name}")
    doc.add_paragraph(f"Date: {datetime.now().strftime('%B %d, %Y')}")
    doc.add_paragraph()

    for i, point in enumerate([
        f"{settings.COMPANY_NAME} is NOT a law firm.",
        "Staff members are NOT licensed attorneys.",
        "Services provided are ADMINISTRATIVE and EDUCATIONAL only.",
        "No legal advice will be given under any circumstances.",
        "We do NOT represent clients in court proceedings.",
        "We do NOT file legal pleadings, motions, or petitions.",
        "Legal services require referral to a licensed attorney.",
        "Administrative assistance does NOT guarantee any legal outcome.",
        "I have the right to hire an attorney at any time.",
        "I understand and agree to the above conditions.",
    ], 1):
        doc.add_paragraph(f"{i}. {point}")

    doc.add_paragraph()
    doc.add_paragraph(
        f"I, {name}, hereby acknowledge that I have read, understand, and agree to the above disclosures."
    )
    doc.add_paragraph()
    _sig_block(doc, "Client")

    path = _filename("non_attorney_disclosure", data.get('client_id', 'new'))
    doc.save(path)
    return path


def gen_authorization_retrieve_records(data: Dict[str, Any]) -> str:
    doc = _base_doc("Authorization to Retrieve Records")
    name = f"{data.get('first_name','')} {data.get('last_name','')}".strip()

    doc.add_paragraph(
        f"I, {name}, born {data.get('dob','')}, hereby authorize {settings.COMPANY_NAME} and/or "
        "its designated representatives to request and retrieve the following records on my behalf:"
    )
    doc.add_paragraph()

    for record_type in data.get('record_types', ['Criminal history records', 'Court records', 'Disposition records']):
        doc.add_paragraph(f"  ✓ {record_type}")

    doc.add_paragraph()
    doc.add_paragraph(
        "This authorization is valid for 90 days from the date signed. I understand that "
        f"{settings.COMPANY_NAME} is retrieving these records for administrative purposes only."
    )
    doc.add_paragraph()
    doc.add_paragraph(f"Full Legal Name: {name}")
    doc.add_paragraph(f"Date of Birth: {data.get('dob','')}")
    doc.add_paragraph(f"SSN (Last 4): XXX-XX-{data.get('ssn_last4','****')}")
    doc.add_paragraph(f"State(s) to Search: {data.get('search_states', 'South Carolina')}")
    doc.add_paragraph()
    _sig_block(doc, "Client")

    path = _filename("authorization_records", data.get('client_id', 'new'))
    doc.save(path)
    return path


def gen_employment_explanation_letter(data: Dict[str, Any]) -> str:
    doc = _base_doc("Employment Explanation Letter")
    name = f"{data.get('first_name','')} {data.get('last_name','')}".strip()

    doc.add_paragraph("To Whom It May Concern,")
    doc.add_paragraph()
    doc.add_paragraph(
        f"My name is {name} and I am writing to provide context regarding a matter in my past "
        "as it relates to my employment application."
    )
    doc.add_paragraph()
    doc.add_paragraph(
        f"In {data.get('incident_year','')}, I experienced a situation that resulted in "
        f"{data.get('conviction_brief','a matter on my record')}. I take full responsibility for "
        "my actions and have spent considerable time reflecting, growing, and working toward becoming "
        "a better version of myself."
    )
    doc.add_paragraph()
    doc.add_paragraph(data.get('rehabilitation_statement',
        "Since that time, I have completed [programs/education/training], maintained stable employment, "
        "and demonstrated a consistent commitment to responsible living. I am eager to bring my skills, "
        "work ethic, and dedication to your organization."
    ))
    doc.add_paragraph()
    doc.add_paragraph(
        "I understand that this information may raise questions, and I welcome the opportunity to "
        "discuss it directly. I am confident that my recent record speaks to the person I am today."
    )
    doc.add_paragraph()
    doc.add_paragraph(f"Respectfully,\n\n{name}")
    doc.add_paragraph(f"Phone: {data.get('phone','')}")
    doc.add_paragraph(f"Email: {data.get('email','')}")

    path = _filename("employment_letter", data.get('client_id', 'new'))
    doc.save(path)
    return path


def gen_pardon_support_packet(data: Dict[str, Any]) -> List[str]:
    files = []

    # 1. Cover letter
    doc = _base_doc("Pardon Support Packet — Cover Letter")
    name = f"{data.get('first_name','')} {data.get('last_name','')}".strip()

    doc.add_paragraph(f"To: South Carolina Board of Pardons, Paroles and Community Supervision")
    doc.add_paragraph()
    doc.add_paragraph(f"Re: Pardon Application — {name}")
    doc.add_paragraph()
    doc.add_paragraph(
        f"Enclosed please find the pardon support materials for {name}. This packet has been assembled "
        f"by {settings.COMPANY_NAME} as an administrative service. All materials are provided for "
        "informational and organizational purposes only. No legal representation is implied."
    )
    doc.add_paragraph()
    doc.add_paragraph("PACKET CONTENTS:")
    for item in [
        "Cover Letter (this document)",
        "Client Background Summary",
        "Employment Explanation Letter",
        "Character Reference Template",
        "Supporting Documentation Index",
    ]:
        doc.add_paragraph(f"  □ {item}")

    path1 = _filename("pardon_cover", data.get('client_id', 'new'))
    doc.save(path1)
    files.append(path1)

    # 2. Character reference template
    doc2 = _base_doc("Character Reference Letter Template")
    doc2.add_paragraph("[DATE]")
    doc2.add_paragraph()
    doc2.add_paragraph("To the South Carolina Board of Pardons, Paroles and Community Supervision:")
    doc2.add_paragraph()
    doc2.add_paragraph(
        f"My name is [REFERENCE NAME] and I have known {name} for [NUMBER] years in my capacity as "
        "[RELATIONSHIP: employer / pastor / neighbor / community leader]."
    )
    doc2.add_paragraph()
    doc2.add_paragraph(
        f"During the time I have known {name}, I have observed [describe character, work ethic, "
        "community involvement, personal growth]. I believe that [he/she/they] has demonstrated "
        "genuine rehabilitation and is committed to making positive contributions to society."
    )
    doc2.add_paragraph()
    doc2.add_paragraph(
        "I write this letter in full support of this pardon application and am available to be "
        "contacted for further information."
    )
    doc2.add_paragraph()
    doc2.add_paragraph("Respectfully,\n\n[REFERENCE NAME]\n[TITLE]\n[PHONE]\n[EMAIL]")

    path2 = _filename("character_reference_template", data.get('client_id', 'new'))
    doc2.save(path2)
    files.append(path2)

    return files


def gen_reentry_resume(data: Dict[str, Any]) -> str:
    doc = _base_doc(f"Resume — {data.get('first_name','')} {data.get('last_name','')}")
    name = f"{data.get('first_name','')} {data.get('last_name','')}".strip()

    # Contact
    cp = doc.add_paragraph()
    cp.add_run(name).bold = True
    cp.add_run(f"\n{data.get('phone','')}  ·  {data.get('email','')}")
    if data.get('city'):
        cp.add_run(f"  ·  {data.get('city','')}, {data.get('state','SC')}")
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph()
    doc.add_paragraph("OBJECTIVE", style="Heading 2")
    doc.add_paragraph(data.get('objective',
        f"Motivated and dedicated professional seeking opportunities to apply skills and experience "
        "while contributing meaningfully to a team environment."))

    if data.get('skills'):
        doc.add_paragraph("SKILLS", style="Heading 2")
        for skill in data.get('skills', []):
            doc.add_paragraph(f"  • {skill}")

    if data.get('experience'):
        doc.add_paragraph("EXPERIENCE", style="Heading 2")
        for exp in data.get('experience', []):
            doc.add_paragraph(f"{exp.get('title','')} — {exp.get('employer','')} ({exp.get('dates','')})")
            doc.add_paragraph(exp.get('description', ''))

    if data.get('education'):
        doc.add_paragraph("EDUCATION", style="Heading 2")
        for edu in data.get('education', []):
            doc.add_paragraph(f"{edu.get('degree','')} — {edu.get('institution','')} ({edu.get('year','')})")

    if data.get('certifications'):
        doc.add_paragraph("CERTIFICATIONS & TRAINING", style="Heading 2")
        for cert in data.get('certifications', []):
            doc.add_paragraph(f"  • {cert}")

    doc.add_paragraph()
    doc.add_paragraph("References available upon request.", style="Normal")

    path = _filename("resume", data.get('client_id', 'new'))
    doc.save(path)
    return path


# ─────────────────────────────────────────────
# DIVISION 4 — DOCUMENT PREPARATION
# ─────────────────────────────────────────────

def gen_affidavit(data: Dict[str, Any]) -> str:
    doc = _base_doc("Affidavit")
    name = f"{data.get('first_name','')} {data.get('last_name','')}".strip()
    state = data.get('state', 'South Carolina')
    county = data.get('county', '')

    doc.add_paragraph(f"STATE OF {state.upper()}")
    doc.add_paragraph(f"COUNTY OF {county.upper()}")
    doc.add_paragraph()
    doc.add_paragraph(f"I, {name}, being duly sworn, depose and state as follows:")
    doc.add_paragraph()

    for i, statement in enumerate(data.get('statements', ['[Statement of fact]']), 1):
        doc.add_paragraph(f"{i}. {statement}")

    doc.add_paragraph()
    doc.add_paragraph(
        "I declare under penalty of perjury under the laws of the State of South Carolina "
        "that the foregoing is true and correct to the best of my knowledge."
    )
    doc.add_paragraph()
    _sig_block(doc, "Affiant")
    doc.add_paragraph()
    doc.add_paragraph()
    doc.add_paragraph("Subscribed and sworn before me this ____ day of ____________, 20____.")
    doc.add_paragraph()
    doc.add_paragraph("_" * 40)
    doc.add_paragraph("Notary Public — State of South Carolina")
    doc.add_paragraph("My Commission Expires: _______________")

    path = _filename("affidavit", data.get('client_id', 'new'))
    doc.save(path)
    return path


def gen_general_correspondence(data: Dict[str, Any]) -> str:
    doc = _base_doc(data.get('subject', 'Administrative Correspondence'))

    doc.add_paragraph(f"To: {data.get('recipient_name','')}")
    if data.get('recipient_org'):
        doc.add_paragraph(f"Organization: {data.get('recipient_org','')}")
    doc.add_paragraph(f"Re: {data.get('subject','')}")
    doc.add_paragraph()

    for paragraph in data.get('body_paragraphs', ['[Letter body]']):
        doc.add_paragraph(paragraph)

    doc.add_paragraph()
    doc.add_paragraph(f"Respectfully,\n\n{data.get('first_name','')} {data.get('last_name','')}")
    doc.add_paragraph(f"Phone: {data.get('phone','')}")
    doc.add_paragraph(f"Email: {data.get('email','')}")

    path = _filename("correspondence", data.get('client_id', 'new'))
    doc.save(path)
    return path


# ─────────────────────────────────────────────
# DIVISION 5 — JUDGMENT & ASSET RECOVERY
# ─────────────────────────────────────────────

def gen_asset_research_report(data: Dict[str, Any]) -> str:
    doc = _base_doc("Asset Research Report")
    name = f"{data.get('first_name','')} {data.get('last_name','')}".strip()

    doc.add_paragraph(f"Prepared For: {name}")
    doc.add_paragraph(f"Research Type: {data.get('research_type', 'Unclaimed Property')}")
    doc.add_paragraph(f"Research Date: {datetime.now().strftime('%B %d, %Y')}")
    doc.add_paragraph(f"Reference #: CA-{data.get('client_id','')}-{datetime.now().strftime('%Y%m%d')}")
    doc.add_paragraph()

    doc.add_paragraph("EXECUTIVE SUMMARY", style="Heading 2")
    doc.add_paragraph(data.get('summary', 'Research completed. Findings detailed below.'))
    doc.add_paragraph()

    doc.add_paragraph("RESEARCH FINDINGS", style="Heading 2")
    findings = data.get('findings', [])
    if findings:
        for f in findings:
            doc.add_paragraph(f"  • Source: {f.get('source','')}")
            doc.add_paragraph(f"    Description: {f.get('description','')}")
            doc.add_paragraph(f"    Estimated Value: {f.get('value','')}")
            doc.add_paragraph()
    else:
        doc.add_paragraph("No findings at this time. Expanded research may be warranted.")

    doc.add_paragraph("RECOMMENDED NEXT STEPS", style="Heading 2")
    for step in data.get('next_steps', ['Contact our office to discuss recovery options.']):
        doc.add_paragraph(f"  → {step}")

    doc.add_paragraph()
    doc.add_paragraph(
        f"DISCLAIMER: This report is prepared for informational purposes only by {settings.COMPANY_NAME}. "
        "It does not constitute legal advice. Recovery of assets may require legal action through a licensed attorney."
    )

    path = _filename("asset_report", data.get('client_id', 'new'))
    doc.save(path)
    return path


def gen_surplus_funds_packet(data: Dict[str, Any]) -> str:
    doc = _base_doc("Surplus Funds Research Packet")
    name = f"{data.get('first_name','')} {data.get('last_name','')}".strip()

    doc.add_paragraph(f"Client: {name}")
    doc.add_paragraph(f"Property Address: {data.get('property_address','')}")
    doc.add_paragraph(f"Sale Date: {data.get('sale_date','')}")
    doc.add_paragraph(f"Estimated Surplus: ${data.get('estimated_surplus', 0):.2f}")
    doc.add_paragraph()

    doc.add_paragraph("PROPERTY DETAILS", style="Heading 2")
    doc.add_paragraph(f"Foreclosure/Tax Sale Date: {data.get('sale_date','')}")
    doc.add_paragraph(f"Winning Bid: ${data.get('winning_bid', 0):.2f}")
    doc.add_paragraph(f"Debt Owed: ${data.get('debt_owed', 0):.2f}")
    doc.add_paragraph(f"Estimated Surplus Funds: ${data.get('estimated_surplus', 0):.2f}")
    doc.add_paragraph()

    doc.add_paragraph("CLAIM PROCESS OVERVIEW", style="Heading 2")
    for step in [
        "Verify surplus funds are available with the county clerk",
        "Gather documentation: proof of ownership, ID, liens",
        "Submit claim (attorney may be required for court filing)",
        "Await disbursement — typically 30–90 days",
    ]:
        doc.add_paragraph(f"  {step}")

    doc.add_paragraph()
    doc.add_paragraph(
        f"NOTE: {settings.COMPANY_NAME} provides research and administrative assistance only. "
        "Filing claims in court requires a licensed attorney."
    )

    path = _filename("surplus_funds", data.get('client_id', 'new'))
    doc.save(path)
    return path


# ─────────────────────────────────────────────
# DIVISION 6 — BUSINESS CONSULTING
# ─────────────────────────────────────────────

def gen_llc_checklist(data: Dict[str, Any]) -> str:
    doc = _base_doc("LLC Formation Checklist")
    name = f"{data.get('first_name','')} {data.get('last_name','')}".strip()

    doc.add_paragraph(f"Prepared For: {name}")
    doc.add_paragraph(f"Business Name: {data.get('business_name','[Business Name]')}")
    doc.add_paragraph(f"State of Formation: {data.get('formation_state', 'South Carolina')}")
    doc.add_paragraph()

    doc.add_paragraph("PRE-FORMATION", style="Heading 2")
    for item in [
        "Choose and verify business name availability (SC Secretary of State)",
        "Determine LLC structure (single-member vs. multi-member)",
        "Choose registered agent (person or service)",
        "Decide on management structure (member-managed vs. manager-managed)",
    ]:
        doc.add_paragraph(f"  ☐ {item}")

    doc.add_paragraph("FORMATION DOCUMENTS", style="Heading 2")
    for item in [
        "File Articles of Organization with SC Secretary of State ($110 fee)",
        "Draft and sign Operating Agreement",
        "Obtain EIN from IRS (free at IRS.gov)",
    ]:
        doc.add_paragraph(f"  ☐ {item}")

    doc.add_paragraph("POST-FORMATION", style="Heading 2")
    for item in [
        "Open business bank account",
        "Register for SC state taxes (SC DOR)",
        "Obtain required business licenses",
        "Set up business accounting system",
        "Get business insurance",
        "Build company website and Google Business Profile",
    ]:
        doc.add_paragraph(f"  ☐ {item}")

    path = _filename("llc_checklist", data.get('client_id', 'new'))
    doc.save(path)
    return path


def gen_business_plan(data: Dict[str, Any]) -> str:
    doc = _base_doc(f"Business Plan — {data.get('business_name','[Business Name]')}")

    doc.add_paragraph(f"Owner: {data.get('first_name','')} {data.get('last_name','')}")
    doc.add_paragraph(f"Date: {datetime.now().strftime('%B %d, %Y')}")
    doc.add_paragraph()

    sections = [
        ("Executive Summary",      data.get('executive_summary',      '[Summarize your business in 2-3 sentences]')),
        ("Business Description",   data.get('business_description',   '[Describe what your business does]')),
        ("Products & Services",    data.get('products_services',      '[List what you sell or offer]')),
        ("Target Market",          data.get('target_market',          '[Describe your ideal customer]')),
        ("Competitive Analysis",   data.get('competitive_analysis',   '[Who are your competitors and how do you differ?]')),
        ("Marketing Strategy",     data.get('marketing_strategy',     '[How will you reach customers?]')),
        ("Operations Plan",        data.get('operations_plan',        '[How will the business run day to day?]')),
        ("Financial Projections",  data.get('financial_projections',  '[Revenue, expenses, profit projections]')),
        ("Funding Requirements",   data.get('funding_requirements',   '[Capital needed and how it will be used]')),
    ]

    for title, content in sections:
        doc.add_paragraph(title, style="Heading 2")
        doc.add_paragraph(content)
        doc.add_paragraph()

    path = _filename("business_plan", data.get('client_id', 'new'))
    doc.save(path)
    return path


def gen_consultation_report(data: Dict[str, Any]) -> str:
    doc = _base_doc("Consultation Report")
    name = f"{data.get('first_name','')} {data.get('last_name','')}".strip()

    doc.add_paragraph(f"Client: {name}")
    doc.add_paragraph(f"Session Date: {data.get('session_date', datetime.now().strftime('%B %d, %Y'))}")
    doc.add_paragraph(f"Session Type: {data.get('session_type', 'Business Consulting')}")
    doc.add_paragraph(f"Duration: {data.get('duration', '60')} minutes")
    doc.add_paragraph()

    doc.add_paragraph("TOPICS DISCUSSED", style="Heading 2")
    for topic in data.get('topics', ['[Topic]']):
        doc.add_paragraph(f"  • {topic}")

    doc.add_paragraph("KEY FINDINGS & RECOMMENDATIONS", style="Heading 2")
    for rec in data.get('recommendations', ['[Recommendation]']):
        doc.add_paragraph(f"  → {rec}")

    doc.add_paragraph("ACTION ITEMS", style="Heading 2")
    for i, action in enumerate(data.get('action_items', []), 1):
        doc.add_paragraph(f"  {i}. {action.get('task','')} — Due: {action.get('due_date','TBD')}")

    doc.add_paragraph(f"\nNext Session: {data.get('next_session', 'To be scheduled')}")

    path = _filename("consultation_report", data.get('client_id', 'new'))
    doc.save(path)
    return path


# ─────────────────────────────────────────────
# DISPATCH — maps document_type → generator
# ─────────────────────────────────────────────

DOCUMENT_REGISTRY = {
    # Notary
    "notary_service_agreement":    gen_notary_service_agreement,
    "notary_invoice":              gen_notary_invoice,
    # Credit
    "credit_services_agreement":   gen_credit_services_agreement,
    "dispute_letter":              gen_dispute_letter,
    "debt_validation_letter":      gen_debt_validation_letter,
    "goodwill_letter":             gen_goodwill_letter,
    # Reentry
    "reentry_service_agreement":   gen_reentry_service_agreement,
    "non_attorney_disclosure":     gen_non_attorney_disclosure,
    "authorization_records":       gen_authorization_retrieve_records,
    "employment_explanation":      gen_employment_explanation_letter,
    "pardon_support_packet":       gen_pardon_support_packet,
    "reentry_resume":              gen_reentry_resume,
    # Document Prep
    "affidavit":                   gen_affidavit,
    "general_correspondence":      gen_general_correspondence,
    # Asset Recovery
    "asset_research_report":       gen_asset_research_report,
    "surplus_funds_packet":        gen_surplus_funds_packet,
    # Business
    "llc_checklist":               gen_llc_checklist,
    "business_plan":               gen_business_plan,
    "consultation_report":         gen_consultation_report,
}

def generate(document_type: str, data: Dict[str, Any]) -> Union[str, List[str]]:
    fn = DOCUMENT_REGISTRY.get(document_type)
    if not fn:
        raise ValueError(f"Unknown document type: '{document_type}'. Available: {list(DOCUMENT_REGISTRY.keys())}")
    return fn(data)

def available_documents() -> Dict[str, List[str]]:
    return {
        "notary":          ["notary_service_agreement", "notary_invoice"],
        "credit":          ["credit_services_agreement", "dispute_letter", "debt_validation_letter", "goodwill_letter"],
        "reentry":         ["reentry_service_agreement", "non_attorney_disclosure", "authorization_records",
                            "employment_explanation", "pardon_support_packet", "reentry_resume"],
        "document_prep":   ["affidavit", "general_correspondence"],
        "asset_recovery":  ["asset_research_report", "surplus_funds_packet"],
        "business":        ["llc_checklist", "business_plan", "consultation_report"],
    }
